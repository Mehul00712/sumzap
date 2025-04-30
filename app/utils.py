import os
import random
import string
from datetime import datetime, timedelta
import PyPDF2
from docx import Document as DocxDocument
import openai
import io
from werkzeug.utils import secure_filename
from app import db
from app.models import TempEmail

def extract_text_from_pdf(pdf_path):
    """Extract text content from a PDF file."""
    text = ""
    with open(pdf_path, 'rb') as file:
        pdf_reader = PyPDF2.PdfReader(file)
        for page_num in range(len(pdf_reader.pages)):
            page = pdf_reader.pages[page_num]
            text += page.extract_text()
    return text

def extract_text_from_docx(docx_path):
    """Extract text content from a DOCX file."""
    doc = DocxDocument(docx_path)
    return " ".join([para.text for para in doc.paragraphs])

def summarize_text(text, api_key):
    """Use OpenAI API to summarize text."""
    if not api_key:
        return "API key not configured. Please add your OpenAI API key to use the summarization feature."
    
    openai.api_key = api_key
    
    try:
        # Limit text to avoid token limits
        if len(text) > 4000:
            text = text[:4000] + "..."
        
        response = openai.chat.completions.create(
            model="gpt-3.5-turbo",
            messages=[
                {"role": "system", "content": "You are a helpful assistant that summarizes documents concisely."},
                {"role": "user", "content": f"Please summarize the following text in a comprehensive way, capturing the key points and main ideas: {text}"}
            ],
            max_tokens=500
        )
        return response.choices[0].message.content
    except Exception as e:
        return f"Error during summarization: {str(e)}"

def convert_pdf_to_docx(pdf_path):
    """Convert PDF to DOCX format."""
    # Extract text from PDF
    text = extract_text_from_pdf(pdf_path)
    
    # Create a new Word document
    doc = DocxDocument()
    doc.add_paragraph(text)
    
    # Create output filename
    filename = os.path.basename(pdf_path)
    base_name = os.path.splitext(filename)[0]
    docx_path = os.path.join(os.path.dirname(pdf_path), f"{base_name}.docx")
    
    # Save the Word document
    doc.save(docx_path)
    
    return docx_path

def generate_temp_email():
    """Generate a temporary email address."""
    # This is just a simulation - in a real app, you would integrate with a temporary email service API
    username = ''.join(random.choices(string.ascii_lowercase + string.digits, k=10))
    domain = random.choice(['tempmail.org', 'temporarymail.com', 'disposable.com'])
    email = f"{username}@{domain}"
    
    # Generate a random password
    password = ''.join(random.choices(string.ascii_letters + string.digits + string.punctuation, k=12))
    
    # Save to database
    temp_email = TempEmail(email=email, password=password)
    db.session.add(temp_email)
    db.session.commit()
    
    return email, password

def clean_expired_emails():
    """Remove expired temporary emails from the database."""
    now = datetime.utcnow()
    expired_emails = TempEmail.query.filter(TempEmail.expires_at <= now).all()
    
    for email in expired_emails:
        db.session.delete(email)
    
    db.session.commit() 
